import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION

# ---- Load data ----
df = pd.read_excel('copaw_gateway_tasks_1.xlsx')

mod_counts = df['模块'].value_counts().sort_index()          # bar: per module
status_counts = df['状态'].value_counts().reindex(['已完成','待开发','需外部'])
prio_counts = df['优先级'].value_counts().reindex(['P0','P1','P2'])

# cumulative completed over task 序号 for line chart
d = df.sort_values('序号')
cum_done = (d['状态'] == '已完成').cumsum()

BRAND = {
    'primary': RGBColor(0, 51, 102),
    'secondary': RGBColor(0, 153, 204),
    'accent': RGBColor(255, 102, 0),
    'green': RGBColor(76, 175, 80),
}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_title_underline(slide, text, top, left=Inches(0.6), width=Inches(12.1)):
    tb = slide.shapes.add_textbox(left, top, width, Inches(0.6))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = BRAND['primary']

def add_subtitle(slide, text, top=Inches(1.05)):
    sub = slide.shapes.add_textbox(Inches(0.6), top, Inches(12.1), Inches(0.5))
    p = sub.text_frame.paragraphs[0]; r = p.add_run(); r.text = text
    r.font.size = Pt(14); r.font.color.rgb = RGBColor(90, 90, 90)

# ============ SLIDE 1: Title ============
s = prs.slides.add_slide(prs.slide_layouts[6])
bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12))
bar.fill.solid(); bar.fill.fore_color.rgb = BRAND['secondary']; bar.line.fill.background()

tb = s.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(10.9), Inches(1.8))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; r = p.add_run(); r.text = "COPaw Gateway 任务进度看板"
r.font.size = Pt(44); r.font.bold = True; r.font.color.rgb = BRAND['primary']
p2 = tf.add_paragraph(); r2 = p2.add_run()
r2.text = "基于 copaw_gateway_tasks_1.xlsx 生成的可交互柱状图与折线图"
r2.font.size = Pt(20); r2.font.color.rgb = BRAND['secondary']

tb2 = s.shapes.add_textbox(Inches(1.2), Inches(4.3), Inches(10.9), Inches(2.0))
tf2 = tb2.text_frame; tf2.word_wrap = True
stats = [
    f"任务总数 {len(df)}　|　模块数 {df['模块'].nunique()}　|　优先级 P0 {prio_counts['P0']} / P1 {prio_counts['P1']} / P2 {prio_counts['P2']}",
    f"已完成 {status_counts['已完成']}　待开发 {status_counts['待开发']}　需外部 {status_counts['需外部']}",
]
for i, line in enumerate(stats):
    p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
    r = p.add_run(); r.text = line
    r.font.size = Pt(18)
    r.font.color.rgb = BRAND['primary'] if i == 0 else RGBColor(80, 80, 80)
    if i == 0: r.font.bold = True

# ============ SLIDE 2: Bar chart - tasks per module ============
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title_underline(s, "各模块任务数量（柱状图）", Inches(0.35))
add_subtitle(s, "按模块统计所有任务的数量分布 —— 原生 Office 图表，可点击/编辑源数据实现交互")
chart_data = CategoryChartData()
chart_data.categories = list(mod_counts.index)
chart_data.add_series('任务数', list(mod_counts.values))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
    Inches(0.6), Inches(1.7), Inches(12.1), Inches(5.4), chart_data)
chart = gf.chart
chart.has_legend = False
chart.has_title = True
chart.chart_title.text_frame.text = "各模块任务数量"
plot = chart.plots[0]
plot.has_data_labels = True
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot.data_labels.font.size = Pt(12); plot.data_labels.font.bold = True
colors = [BRAND['secondary'], BRAND['accent'], BRAND['green'], BRAND['primary']]
series = plot.series[0]
for idx, pt in enumerate(series.points):
    pt.format.fill.solid(); pt.format.fill.fore_color.rgb = colors[idx % len(colors)]
va = chart.value_axis
va.has_major_gridlines = True
va.maximum_scale = 9
va.tick_labels.font.size = Pt(12)
chart.category_axis.tick_labels.font.size = Pt(12)

# ============ SLIDE 3: Line chart - cumulative completion ============
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title_underline(s, "任务完成进度（折线图）", Inches(0.35))
add_subtitle(s, "按任务编号累计的已完成数量 vs 累计全部任务 —— 原生 Office 图表，可交互查看")
line_data = CategoryChartData()
line_data.categories = [f"#{i}" for i in d['序号']]
line_data.add_series('累计已完成', list(cum_done))
line_data.add_series('累计全部任务', list(range(1, len(d) + 1)))
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS,
    Inches(0.6), Inches(1.7), Inches(12.1), Inches(5.4), line_data)
lc = gf.chart
lc.has_title = True
lc.chart_title.text_frame.text = "累计完成进度曲线"
lc.legend.position = XL_LEGEND_POSITION.BOTTOM
lc.legend.include_in_layout = False
lc.legend.font.size = Pt(12)
lplot = lc.plots[0]
from pptx.enum.chart import XL_MARKER_STYLE
for ser in lplot.series:
    ser.marker.style = XL_MARKER_STYLE.CIRCLE; ser.marker.size = 6
lplot.has_data_labels = True
lplot.data_labels.show_value = True
lplot.data_labels.font.size = Pt(9)
lva = lc.value_axis
lva.has_major_gridlines = True
lva.maximum_scale = len(d) + 1
lva.major_gridlines.format.line.color.rgb = RGBColor(217, 217, 217)
lc.category_axis.tick_labels.font.size = Pt(9)

# ============ SLIDE 4: Status & priority overview ============
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title_underline(s, "任务状态与优先级概览", Inches(0.35))

cd2 = CategoryChartData()
cd2.categories = list(status_counts.index)
cd2.add_series('数量', list(status_counts.values))
gf2 = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
    Inches(0.6), Inches(1.6), Inches(6.0), Inches(5.2), cd2)
c2 = gf2.chart
c2.has_legend = False; c2.has_title = True
c2.chart_title.text_frame.text = "任务状态分布"
p2_ = c2.plots[0]; p2_.has_data_labels = True
p2_.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
for idx, pt in enumerate(p2_.series[0].points):
    pt.format.fill.solid(); pt.format.fill.fore_color.rgb = colors[idx % len(colors)]

cd3 = CategoryChartData()
cd3.categories = list(prio_counts.index)
cd3.add_series('数量', list(prio_counts.values))
gf3 = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
    Inches(6.9), Inches(1.6), Inches(6.0), Inches(5.2), cd3)
c3 = gf3.chart
c3.has_legend = False; c3.has_title = True
c3.chart_title.text_frame.text = "优先级分布"
p3_ = c3.plots[0]; p3_.has_data_labels = True
p3_.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
for idx, pt in enumerate(p3_.series[0].points):
    pt.format.fill.solid(); pt.format.fill.fore_color.rgb = colors[idx % len(colors)]

prs.save('copaw_gateway_dashboard.pptx')
print("Saved copaw_gateway_dashboard.pptx | slides:", len(prs.slides._sldIdLst))
